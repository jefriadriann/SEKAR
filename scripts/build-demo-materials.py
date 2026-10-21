#!/usr/bin/env python3
"""Membangun berkas materi demo SEKAR (SEMUA SIMULASI).

    python3 scripts/build-demo-materials.py

Keluaran:
  public/demo/sgo/*.xlsx         worksheet SGo per area (Aktivitas & Supervisi, Pengendalian)
  public/demo/ketentuan/*.pdf    ketentuan demo per area (PDG, PADGI, Juknis, SE Intern)
  public/demo/video/*.mp4|webm   video panduan pencatatan aset + video pengenalan SGo
  data/base/demo-materials.json  manifest yang dibaca scripts/generate-dummy-data.mjs

Butuh: openpyxl, reportlab, Pillow, dan ffmpeg (libx264). Hasil deterministik.
"""
import json
import os
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from demo_content import (  # noqa: E402
    ACTIVITIES,
    AREAS,
    ASSET_VIDEOS,
    KETENTUAN_TOPIC,
    KETENTUAN_TYPES,
    SGO_INTRO_STEPS,
    SLUG,
)

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
PUB = os.path.join(ROOT, "public", "demo")
NAVY = "13235A"
BLUE = "1D4F9E"


def out(*parts):
    p = os.path.join(PUB, *parts)
    os.makedirs(os.path.dirname(p), exist_ok=True)
    return p


# --------------------------------------------------------------------------- Excel
THIN = Side(style="thin", color="B7C3D9")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor=BLUE)
BAND_FILL = PatternFill("solid", fgColor="EEF2F9")
INPUT_FILL = PatternFill("solid", fgColor="FFF8E1")
WRAP = Alignment(wrap_text=True, vertical="top")


def title_block(ws, title, area, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    ws["A2"] = "DOKUMEN SIMULASI — contoh format untuk demo SEKAR, bukan worksheet resmi."
    ws["A2"].font = Font(italic=True, size=9, color="C4323B")
    meta = [("Satuan kerja", "KPw Simulasi __"), ("Area pemeriksaan", area), ("Periode", "Semester II 2026"), ("Disusun oleh", "PIC Simulasi"), ("Disupervisi oleh", "Kepala Unit Simulasi")]
    for i, (k, v) in enumerate(meta):
        r = 4 + i
        ws.cell(row=r, column=1, value=k).font = Font(bold=True, color=NAVY)
        c = ws.cell(row=r, column=3, value=v)
        c.fill = INPUT_FILL
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    return 10


def header_row(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BOX
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[row].height = 34


def body_rows(ws, start, rows, input_cols):
    for r_i, row in enumerate(rows):
        r = start + r_i
        for c_i, v in enumerate(row, start=1):
            c = ws.cell(row=r, column=c_i, value=v)
            c.alignment = WRAP
            c.border = BOX
            if c_i in input_cols:
                c.fill = INPUT_FILL
            elif r_i % 2:
                c.fill = BAND_FILL
        ws.row_dimensions[r].height = 42


def add_list(ws, col_letter, first, last, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}{first}:{col_letter}{last}")


def guide_sheet(wb, lines):
    ws = wb.create_sheet("Petunjuk")
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 110
    ws["B1"] = "Petunjuk Pengisian (SIMULASI)"
    ws["B1"].font = Font(bold=True, size=13, color=NAVY)
    for i, line in enumerate(lines, start=3):
        ws.cell(row=i, column=1, value=i - 2).font = Font(bold=True, color=BLUE)
        ws.cell(row=i, column=2, value=line).alignment = WRAP


def ws_aktivitas(area):
    acts = ACTIVITIES[area]
    headers = ["No", "Proses", "Aktivitas", "Output", "Frekuensi", "PIC Pelaksana", "Dilaksanakan?", "Bukti Pelaksanaan", "Tanggal Supervisi", "Disupervisi oleh", "Hasil Supervisi", "Catatan Supervisi"]
    widths = [5, 20, 38, 24, 14, 16, 14, 26, 15, 18, 18, 34]
    rows = []
    for i, (proses, akt, output, frek, *_rest) in enumerate(acts):
        done = "Ya" if i % 4 != 2 else "Tidak"
        hasil = "Memadai" if done == "Ya" and i % 3 != 1 else "Perlu perbaikan"
        catatan = "Bukti lengkap dan tepat waktu." if hasil == "Memadai" else "Lengkapi bukti dan jadwalkan ulang supervisi."
        rows.append([i + 1, proses, akt, output, frek, f"PIC Simulasi {i % 4 + 1}", done, f"{output} (scan)", f"2026-0{7 + i % 3}-1{i % 9}", "Kepala Unit Simulasi", hasil, catatan])
    wb = Workbook()
    ws = wb.active
    ws.title = "Aktivitas & Supervisi"
    start = title_block(ws, f"WORKSHEET AKTIVITAS DAN SUPERVISI — AREA {area.upper()}", area, len(headers))
    header_row(ws, start, headers, widths)
    body_rows(ws, start + 1, rows, input_cols={7, 8, 9, 10, 11, 12})
    last = start + len(rows)
    add_list(ws, "G", start + 1, last + 20, ["Ya", "Tidak", "Sebagian"])
    add_list(ws, "K", start + 1, last + 20, ["Memadai", "Perlu perbaikan", "Belum disupervisi"])
    ws.freeze_panes = ws.cell(row=start + 1, column=3)

    rk = wb.create_sheet("Rekap")
    rk["A1"] = f"Rekap Aktivitas & Supervisi — {area}"
    rk["A1"].font = Font(bold=True, size=13, color=NAVY)
    ref = f"'Aktivitas & Supervisi'!"
    items = [
        ("Jumlah aktivitas", f"=COUNTA({ref}C{start + 1}:C{last})"),
        ("Dilaksanakan", f'=COUNTIF({ref}G{start + 1}:G{last},"Ya")'),
        ("Tidak dilaksanakan", f'=COUNTIF({ref}G{start + 1}:G{last},"Tidak")'),
        ("Supervisi memadai", f'=COUNTIF({ref}K{start + 1}:K{last},"Memadai")'),
        ("Perlu perbaikan", f'=COUNTIF({ref}K{start + 1}:K{last},"Perlu perbaikan")'),
    ]
    rk.column_dimensions["A"].width = 28
    for i, (k, f) in enumerate(items, start=3):
        rk.cell(row=i, column=1, value=k).font = Font(bold=True)
        rk.cell(row=i, column=2, value=f)
    guide_sheet(wb, [
        "Isi identitas satuan kerja dan periode pada bagian atas lembar Aktivitas & Supervisi.",
        "Untuk setiap aktivitas, pilih 'Dilaksanakan?' dari daftar (Ya/Tidak/Sebagian).",
        "Cantumkan bukti pelaksanaan (nama berkas / nomor dokumen).",
        "Pejabat supervisi mengisi tanggal, nama, hasil supervisi, dan catatan.",
        "Kolom berwarna kuning adalah kolom isian. Lembar Rekap terhitung otomatis.",
        "Simpan worksheet sebagai bukti persiapan pemeriksaan (simulasi).",
    ])
    path = out("sgo", f"{SLUG[area]}-ws-aktivitas-supervisi.xlsx")
    wb.save(path)
    return headers, rows


def ws_pengendalian(area):
    acts = ACTIVITIES[area]
    headers = ["No", "Proses", "Risiko", "Pengendalian Kunci", "Jenis", "Frekuensi", "Prosedur Pengujian", "Jumlah Sampel", "Hasil Pengujian", "Kesimpulan", "Rekomendasi", "Referensi Ketentuan"]
    widths = [5, 20, 32, 36, 12, 14, 34, 10, 30, 14, 32, 22]
    rows = []
    for i, (proses, _akt, _out, frek, risiko, kontrol, jenis) in enumerate(acts):
        efektif = i % 3 != 1
        rows.append([
            i + 1, proses, risiko, kontrol, jenis, frek,
            f"Ambil sampel dan uji apakah '{kontrol.lower()}' dijalankan dan terdokumentasi.",
            5 + (i * 3) % 11,
            "Seluruh sampel sesuai." if efektif else f"{1 + i % 3} sampel tidak didukung bukti.",
            "Efektif" if efektif else "Tidak efektif",
            "-" if efektif else "Perkuat dokumentasi pengendalian dan lakukan reviu berjenjang.",
            f"PADGI SIM/{SLUG[area].upper()[:6]} Pasal {i + 3}",
        ])
    wb = Workbook()
    ws = wb.active
    ws.title = "Pengendalian"
    start = title_block(ws, f"WORKSHEET PENGENDALIAN — AREA {area.upper()}", area, len(headers))
    header_row(ws, start, headers, widths)
    body_rows(ws, start + 1, rows, input_cols={8, 9, 10, 11})
    last = start + len(rows)
    add_list(ws, "E", start + 1, last + 20, ["Preventif", "Detektif", "Korektif"])
    add_list(ws, "J", start + 1, last + 20, ["Efektif", "Tidak efektif", "Belum diuji"])
    ws.freeze_panes = ws.cell(row=start + 1, column=3)

    mr = wb.create_sheet("Matriks Risiko")
    mr["A1"] = f"Matriks Risiko (SIMULASI) — {area}"
    mr["A1"].font = Font(bold=True, size=13, color=NAVY)
    mh = ["No", "Risiko", "Kemungkinan (1-5)", "Dampak (1-5)", "Skor", "Level"]
    for i, h in enumerate(mh, start=1):
        c = mr.cell(row=3, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.border = BOX
    mr.column_dimensions["B"].width = 44
    for c in "CDEF":
        mr.column_dimensions[c].width = 16
    for i, (_p, _a, _o, _f, risiko, *_r) in enumerate(acts):
        r = 4 + i
        mr.cell(row=r, column=1, value=i + 1)
        mr.cell(row=r, column=2, value=risiko)
        mr.cell(row=r, column=3, value=1 + (i * 2) % 5)
        mr.cell(row=r, column=4, value=2 + i % 4)
        mr.cell(row=r, column=5, value=f"=C{r}*D{r}")
        mr.cell(row=r, column=6, value=f'=IF(E{r}>=12,"Tinggi",IF(E{r}>=6,"Sedang","Rendah"))')
    guide_sheet(wb, [
        "Lembar Pengendalian memetakan risiko utama tiap proses dan pengendalian kuncinya.",
        "Tentukan jenis pengendalian (Preventif/Detektif/Korektif) dan frekuensinya.",
        "Lakukan prosedur pengujian atas sampel lalu catat hasilnya.",
        "Pilih kesimpulan Efektif / Tidak efektif; isi rekomendasi bila tidak efektif.",
        "Cantumkan referensi ketentuan (lihat menu Ketentuan di SEKAR).",
        "Lembar Matriks Risiko menghitung skor = kemungkinan x dampak secara otomatis.",
    ])
    path = out("sgo", f"{SLUG[area]}-ws-pengendalian.xlsx")
    wb.save(path)
    return headers, rows


# --------------------------------------------------------------------------- PDF
def watermark(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica-Bold", 54)
    canvas.setFillColor(colors.Color(0.77, 0.2, 0.23, alpha=0.08))
    canvas.translate(A4[0] / 2, A4[1] / 2)
    canvas.rotate(35)
    canvas.drawCentredString(0, 0, "SIMULASI")
    canvas.restoreState()
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#6b7280"))
    canvas.drawString(20 * mm, 12 * mm, "Dokumen demo SEKAR — bukan ketentuan resmi Bank Indonesia.")
    canvas.drawRightString(A4[0] - 20 * mm, 12 * mm, f"Halaman {doc.page}")
    canvas.restoreState()


def ketentuan_pdf(area, i_area, subkind, label, jenis):
    topic = KETENTUAN_TOPIC[area]
    number = f"No. {i_area + 1:02d}/SIM/{label.replace(' ', '').upper()}/2025"
    title = f"{jenis} (Simulasi) {number} tentang {topic}"
    ss = getSampleStyleSheet()
    h = ParagraphStyle("h", parent=ss["Heading2"], textColor=colors.HexColor("#13235A"), spaceBefore=10)
    body = ParagraphStyle("b", parent=ss["BodyText"], fontSize=10, leading=14)
    cover = ParagraphStyle("c", parent=ss["Title"], textColor=colors.HexColor("#13235A"), fontSize=17, leading=22)
    acts = ACTIVITIES[area]
    story = [
        Paragraph(jenis.upper() + " (SIMULASI)", ParagraphStyle("k", parent=body, textColor=colors.HexColor("#C4323B"), alignment=1)),
        Spacer(1, 4),
        Paragraph(f"{number}<br/>tentang<br/>{topic}", cover),
        Spacer(1, 8),
        Paragraph("Dokumen ini dibuat untuk demonstrasi aplikasi SEKAR. Nomor, pasal, dan isi seluruhnya fiktif.", ParagraphStyle("n", parent=body, alignment=1, textColor=colors.HexColor("#6b7280"))),
        Spacer(1, 10),
        Paragraph("BAB I — KETENTUAN UMUM", h),
        Paragraph(f"<b>Pasal 1.</b> Dalam {jenis} ini yang dimaksud dengan area <i>{area}</i> adalah seluruh proses {topic.lower()} pada satuan kerja kantor perwakilan dan departemen.", body),
        Paragraph("<b>Pasal 2.</b> Pengaturan ini bertujuan memberikan acuan baku bagi pelaksana, pejabat supervisi, dan pemeriksa intern.", body),
        Paragraph("BAB II — RUANG LINGKUP", h),
    ]
    for j, (proses, akt, output, frek, risiko, kontrol, _jenis) in enumerate(acts):
        story.append(Paragraph(f"<b>Pasal {j + 3}.</b> {proses}: satuan kerja wajib {akt[0].lower() + akt[1:]} dengan frekuensi {frek.lower()} dan menghasilkan {output.lower()}. Untuk memitigasi risiko <i>{risiko.lower()}</i>, diterapkan pengendalian berupa {kontrol[0].lower() + kontrol[1:]}.", body))
    story += [
        Paragraph("BAB III — MATRIKS KEWAJIBAN", h),
        Spacer(1, 4),
    ]
    table = [["No", "Proses", "Output", "Frekuensi"]] + [[str(j + 1), a[0], a[2], a[3]] for j, a in enumerate(acts)]
    t = Table(table, colWidths=[12 * mm, 52 * mm, 62 * mm, 32 * mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1D4F9E")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B7C3D9")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF2F9")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story += [t, Paragraph("BAB IV — PENUTUP", h)]
    if subkind == "juknis":
        story.append(Paragraph("Petunjuk teknis ini dilengkapi worksheet Aktivitas &amp; Supervisi dan worksheet Pengendalian pada menu SGo.", body))
    elif subkind == "se_intern":
        story.append(Paragraph("Surat edaran ini berlaku sejak tanggal ditetapkan dan disosialisasikan kepada seluruh satuan kerja.", body))
    story.append(Paragraph(f"Ditetapkan di Jakarta (simulasi), 2 Januari 2025. {jenis} ini bersifat contoh dan tidak memiliki kekuatan mengikat.", body))
    path = out("ketentuan", f"{SLUG[area]}-{subkind.replace('_', '-')}.pdf")
    SimpleDocTemplate(path, pagesize=A4, title=title, author="SEKAR Demo (Simulasi)", leftMargin=20 * mm, rightMargin=20 * mm, topMargin=20 * mm, bottomMargin=20 * mm).build(story, onFirstPage=watermark, onLaterPages=watermark)
    return number, title


# --------------------------------------------------------------------------- Video
W, H, FPS = 960, 540, 20
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_B = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"


def font(size, bold=False):
    return ImageFont.truetype(FONT_B if bold else FONT, size)


def wrap(draw, text, fnt, width):
    words, lines, cur = text.split(), [], ""
    for w in words:
        t = (cur + " " + w).strip()
        if draw.textlength(t, font=fnt) <= width:
            cur = t
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def background():
    img = Image.new("RGB", (W, H), "#13235a")
    d = ImageDraw.Draw(img)
    for y in range(H):
        t = y / H
        d.line([(0, y), (W, y)], fill=(int(29 - 10 * t), int(79 - 44 * t), int(158 - 68 * t)))
    d.ellipse([W - 260, -160, W + 140, 240], fill=(36, 92, 170))
    d.ellipse([-120, H - 180, 260, H + 200], fill=(26, 60, 128))
    return img


def scene_base(kicker, title, body_text, number, accent, total_steps=0):
    img = background()
    d = ImageDraw.Draw(img)
    # penanda langkah di bagian bawah
    for k in range(total_steps):
        cx = 72 + k * 46
        done = number is not None and k < number
        d.ellipse([cx - 13, 420 - 13, cx + 13, 420 + 13], fill=accent if done else (40, 62, 120), outline=(190, 205, 235))
        d.text((cx, 420), str(k + 1), font=font(13, True), fill="white", anchor="mm")
    d.text((56, 44), kicker.upper(), font=font(15, True), fill=(190, 205, 235))
    d.text((W - 56, 44), "SIMULASI", font=font(14, True), fill=(255, 210, 120), anchor="ra")
    y = 150
    if number is not None:
        d.rounded_rectangle([56, y, 136, y + 80], radius=20, fill=accent)
        d.text((96, y + 40), str(number), font=font(40, True), fill="white", anchor="mm")
        x0 = 164
    else:
        x0 = 56
    for line in wrap(d, title, font(34, True), W - x0 - 60):
        d.text((x0, y), line, font=font(34, True), fill="white")
        y += 46
    y += 14
    for line in wrap(d, body_text, font(22), W - x0 - 60):
        d.text((x0, y), line, font=font(22), fill=(220, 228, 245))
        y += 32
    return img


def render_video(path, kicker, title, subtitle, steps, accent=(196, 50, 59)):
    scenes = [(scene_base(kicker, title, subtitle, None, accent), 3.0)]
    for i, s in enumerate(steps):
        scenes.append((scene_base(kicker, f"Langkah {i + 1} dari {len(steps)}", s, i + 1, accent, len(steps)), 4.0))
    scenes.append((scene_base(kicker, "Selesai", "Video panduan simulasi SEKAR. Lanjutkan dengan mengisi worksheet dan melampirkan bukti.", None, accent), 2.5))
    total = sum(d for _, d in scenes)
    proc = subprocess.Popen(
        ["ffmpeg", "-y", "-loglevel", "error", "-f", "rawvideo", "-pix_fmt", "rgb24", "-s", f"{W}x{H}", "-r", str(FPS), "-i", "-",
         "-c:v", "libx264", "-preset", "slow", "-crf", "30", "-pix_fmt", "yuv420p", "-movflags", "+faststart", "-an", path],
        stdin=subprocess.PIPE,
    )
    elapsed, prev = 0.0, None
    for img, dur in scenes:
        frames = int(dur * FPS)
        for f in range(frames):
            t = f / FPS
            fade = min(1.0, t / 0.5)
            frame = Image.blend(prev, img, fade) if prev is not None and fade < 1 else img.copy()
            d = ImageDraw.Draw(frame)
            p = (elapsed + t) / total
            d.rectangle([0, H - 8, W, H], fill=(20, 35, 80))
            d.rectangle([0, H - 8, int(W * p), H], fill=accent)
            # indikator "putar" berdenyut
            r = 6 + 2 * ((f // 5) % 2)
            d.ellipse([W - 150 - r, H - 40 - r, W - 150 + r, H - 40 + r], fill=accent)
            d.text((W - 134, H - 40), "PANDUAN", font=font(12, True), fill=(220, 228, 245), anchor="lm")
            proc.stdin.write(frame.tobytes())
        elapsed += dur
        prev = img
    proc.stdin.close()
    proc.wait()
    # Salinan WebM (VP9) untuk browser tanpa dekoder H.264 (mis. Chromium open-source).
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-i", path, "-c:v", "libvpx-vp9", "-b:v", "0", "-crf", "42", "-row-mt", "1", "-an", path[:-4] + ".webm"], check=True)
    return round(total)


# --------------------------------------------------------------------------- main
def main():
    manifest = {"areas": [], "asset_videos": [], "sgo_intro": None}
    for i, area in enumerate(AREAS):
        ah, ar = ws_aktivitas(area)
        ph, pr = ws_pengendalian(area)
        entry = {
            "area": area,
            "slug": SLUG[area],
            "worksheets": [
                {"subkind": "ws_aktivitas_supervisi", "title": f"Worksheet Aktivitas & Supervisi {area}", "file_url": f"/demo/sgo/{SLUG[area]}-ws-aktivitas-supervisi.xlsx",
                 "sheets": ["Aktivitas & Supervisi", "Rekap", "Petunjuk"], "columns": ah, "preview": [[str(c) for c in r] for r in ar[:4]]},
                {"subkind": "ws_pengendalian", "title": f"Worksheet Pengendalian {area}", "file_url": f"/demo/sgo/{SLUG[area]}-ws-pengendalian.xlsx",
                 "sheets": ["Pengendalian", "Matriks Risiko", "Petunjuk"], "columns": ph, "preview": [[str(c) for c in r] for r in pr[:4]]},
            ],
            "ketentuan": [],
        }
        for subkind, label, jenis in KETENTUAN_TYPES:
            number, title = ketentuan_pdf(area, i, subkind, label, jenis)
            entry["ketentuan"].append({"subkind": subkind, "label": label, "jenis": jenis, "number": number, "title": title,
                                        "topic": KETENTUAN_TOPIC[area], "file_url": f"/demo/ketentuan/{SLUG[area]}-{subkind.replace('_', '-')}.pdf"})
        manifest["areas"].append(entry)
        print("materi", area)

    for subkind, label, desc, steps in ASSET_VIDEOS:
        name = f"aset-{subkind.replace('_', '-')}.mp4"
        dur = render_video(out("video", name), "Pencatatan Aset", label, desc, steps)
        manifest["asset_videos"].append({"subkind": subkind, "label": label, "description": desc, "steps": steps, "file_url": f"/demo/video/{name}", "duration": dur})
        print("video", name)

    dur = render_video(out("video", "sgo-pengenalan.mp4"), "Pengenalan SGo", "Mengenal SGo dan Ketentuan", "Ruang materi pemeriksaan bersama KPwDN dan DR.", SGO_INTRO_STEPS, accent=(29, 120, 200))
    manifest["sgo_intro"] = {"title": "Video Pengenalan SGo", "steps": SGO_INTRO_STEPS, "file_url": "/demo/video/sgo-pengenalan.mp4", "duration": dur}

    with open(os.path.join(ROOT, "data", "base", "demo-materials.json"), "w") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=2)
        f.write("\n")
    print("manifest data/base/demo-materials.json")


if __name__ == "__main__":
    main()
